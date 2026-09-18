#!/usr/bin/env python3
"""Read figure4.xlsx back and prove it still says exactly what the exports say.

Regenerates the workbook into a temporary file, reopens it with openpyxl and
checks: sheet names and order; every W&B cell on every arm tab against the
arm's own history.csv under results/ at the line data/manifest.json names (200
rows x all 144/150/57 columns, iterated, not sampled); the frozen leading
block's column order; the numeric twins against the strings they mirror; the
derived trailing mean against the figure builder's moving_average; the README
tab's recipe, provenance and the two charts, including which column each chart
series points at. One negative control proves the cell comparison can fail.

Run from this directory (needs openpyxl and PyYAML):
    python -m unittest make_spreadsheet_test -v
"""

from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import build_end_to_end_results_measured as figure
import make_spreadsheet
from build_end_to_end_results_provisional import Point, moving_average

DIFFERING_KEYS = {"old_logps_source", "sampler_is", "eval_every_n_steps"}
# The endpoint labels the published figure prints, as trailing means.
ENDPOINTS = {"Standard": 0.6265625, "TIS": 0.66875, "Zero-TIM": 0.882421875}
SOLVE_NUM = "derived:solve_ratio_num"
MEAN_NUM = "derived:logp_diff_mean_num"
TRAIL = "derived:solve_ratio_trail10"
LEADING = ["display_step", "_step", "_runtime", "_timestamp",
           "rewards/train/solve_ratio", SOLVE_NUM, TRAIL,
           "sampler_trainer/train/logp_diff_mean", MEAN_NUM,
           "sampler_trainer/train/logp_diff_max"]
DERIVED = {"display_step", SOLVE_NUM, TRAIL, MEAN_NUM}
# Which string column each numeric twin mirrors, and which column each chart plots.
TWINS = {SOLVE_NUM: figure.REWARD, MEAN_NUM: figure.MEAN}
CHART_COLUMNS = (MEAN_NUM, TRAIL)
BYTES_COLUMN = "canonical/train/alignment_max_differing_bytes"
RECIPE_WIDTH = 5


def as_written(value: float) -> float:
    """openpyxl stores a float as "%.16g", so a cell keeps 16 significant digits."""
    return float("%.16g" % value)


def history_rows(arm: str, entry: dict) -> tuple[list, list]:
    """Read the arm's W&B export independently of make_spreadsheet, by line number."""
    path = make_spreadsheet.runs_layout.run_dir(arm, entry["run_id"]) / "history.csv"
    with path.open(newline="", encoding="utf-8") as handle:
        lines = list(csv.reader(handle))
    header = lines[0]
    wanted = set(entry["source_csv_lines"])
    rows = [dict(zip(header, line)) for number, line in enumerate(lines, 1) if number in wanted]
    assert len(rows) == figure.STEPS, f"{path.name}: {len(rows)} plotted rows"
    return header, rows


def mismatches(sheet, header: list, rows: list) -> list:
    """Every W&B cell that is not the export's own string (empty stays empty)."""
    columns = {name: index for index, name in enumerate([c.value for c in sheet[1]], 1)}
    problems = []
    for index, record in enumerate(rows):
        for name in header:
            expected = record[name]
            actual = sheet.cell(row=index + 2, column=columns[name]).value
            if (actual is not None) if expected == "" else (actual != expected):
                problems.append(f"row {index + 2} {name}: {actual!r} != {expected!r}")
    return problems


class SpreadsheetTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cls.path = make_spreadsheet.build(Path(cls.directory.name) / "figure4.xlsx")
        cls.book = load_workbook(cls.path)
        cls.manifest = json.loads(
            (make_spreadsheet.DATA / "manifest.json").read_text(encoding="utf-8"))
        cls.exports = {arm: history_rows(arm, cls.manifest["runs"][arm])
                       for arm, _ in make_spreadsheet.ARMS}

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def readme_block(self, first_cell: str) -> list:
        """The rows under the README block whose header row starts with first_cell."""
        sheet = self.book["README"]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        start = next(index for index, row in enumerate(rows) if row[0] == first_cell)
        block = []
        for row in rows[start + 1:]:
            if row[0] is None:
                break
            block.append(row)
        return block

    def test_sheet_names_and_order(self):
        self.assertEqual(self.book.sheetnames, list(make_spreadsheet.SHEETS))

    def test_arm_tabs_are_the_wandb_export_strings(self):
        for arm, label in make_spreadsheet.ARMS:
            header, rows = self.exports[arm]
            sheet = self.book[label]
            with self.subTest(arm=label):
                self.assertEqual(sheet.max_row, figure.STEPS + 1)
                self.assertEqual(sheet.max_column, len(header) + len(DERIVED))
                self.assertEqual({c.value for c in sheet[1]}, set(header) | DERIVED)
                self.assertEqual(mismatches(sheet, header, rows), [])
                for index in range(figure.STEPS):
                    self.assertEqual(sheet.cell(row=index + 2, column=1).value, index + 1)

    def test_leading_columns_are_identity_then_figure4(self):
        for arm, label in make_spreadsheet.ARMS:
            header, _ = self.exports[arm]
            expected = LEADING + ([BYTES_COLUMN] if BYTES_COLUMN in header else [])
            actual = [cell.value for cell in self.book[label][1]][:len(expected)]
            with self.subTest(arm=label):
                self.assertEqual(actual, expected)
                self.assertEqual(self.book[label].freeze_panes,
                                 f"{chr(ord('A') + len(expected))}2")

    def test_numeric_twins_mirror_the_string_columns(self):
        for arm, label in make_spreadsheet.ARMS:
            _, rows = self.exports[arm]
            for twin, source in TWINS.items():
                column = LEADING.index(twin) + 1
                with self.subTest(arm=label, column=twin):
                    for index, record in enumerate(rows):
                        # The twin is the string's own value, as openpyxl stores it.
                        self.assertEqual(
                            self.book[label].cell(row=index + 2, column=column).value,
                            as_written(float(record[source])),
                            f"{label} {twin} step {index + 1}")

    def test_trailing_mean_is_the_figure_builders(self):
        for arm, label in make_spreadsheet.ARMS:
            _, rows = self.exports[arm]
            points = tuple(Point(index + 1, float(row[figure.REWARD]))
                           for index, row in enumerate(rows))
            expected = moving_average(points)
            column = LEADING.index(TRAIL) + 1
            with self.subTest(arm=label):
                for index, point in enumerate(expected):
                    # Same function on the same floats: the only gap allowed is
                    # openpyxl's 16-significant-digit float rendering.
                    self.assertEqual(self.book[label].cell(row=index + 2, column=column).value,
                                     as_written(point.value), f"{label} step {index + 1}")

    def test_trailing_mean_ends_on_the_published_labels(self):
        column = LEADING.index(TRAIL) + 1
        for label, endpoint in ENDPOINTS.items():
            value = self.book[label].cell(row=figure.STEPS + 1, column=column).value
            self.assertEqual(value, endpoint, label)
            self.assertEqual(f"{value * 100:.1f}%",
                             {"Standard": "62.7%", "TIS": "66.9%",
                              "Zero-TIM": "88.2%"}[label])

    def test_readme_recipe_matches_the_manifest(self):
        rows = self.readme_block("config_key")
        configs = [self.manifest["runs"][arm]["config"] for arm, _ in make_spreadsheet.ARMS]
        self.assertEqual([row[0] for row in rows], list(configs[0]))
        self.assertEqual(len(rows), 24)
        differing = set()
        for key, *values, differs in (row[:RECIPE_WIDTH] for row in rows):
            self.assertIn(differs, (True, False))
            expected = [make_spreadsheet.cell(config[key]) for config in configs]
            self.assertEqual(list(map(str, values)), list(map(str, expected)), key)
            if differs:
                differing.add(key)
        self.assertEqual(differing, DIFFERING_KEYS)

    def test_readme_provenance_matches_the_manifest(self):
        rows = {row[0]: row for row in self.readme_block("arm")}
        for arm, label in make_spreadsheet.ARMS:
            entry = self.manifest["runs"][arm]
            row = rows[label]
            with self.subTest(arm=label):
                self.assertEqual(row[1], entry["run_id"])
                self.assertEqual(row[2], make_spreadsheet.WANDB_PROJECT)
                self.assertTrue(entry["run_id"].endswith(row[3]))
                self.assertEqual(row[4], entry["source_history_sha256"])
                self.assertEqual(row[5], entry["source_config_sha256"])
                self.assertEqual((row[6], row[7]), (entry["source_csv_lines"][0],
                                                    entry["source_csv_lines"][-1]))
                self.assertEqual(row[8], self.manifest["source_commit"])
                self.assertEqual(row[9], self.manifest["evidence_grade"])

    def test_readme_holds_the_two_figure4_charts(self):
        charts = self.book["README"]._charts
        self.assertEqual(len(charts), 2)
        for chart, column in zip(charts, CHART_COLUMNS):
            self.assertEqual(len(chart.series), 3)
            # A chart cannot plot text, so both must point at a numeric column.
            letter = get_column_letter(LEADING.index(column) + 1)
            expected = f"!${letter}$2:${letter}${figure.STEPS + 1}"
            for series, (_, label) in zip(chart.series, make_spreadsheet.ARMS):
                reference = series.val.numRef.f
                with self.subTest(chart=column, arm=label):
                    self.assertTrue(reference.endswith(expected), reference)
                    self.assertIn(label, reference)
                    self.assertEqual(self.book[label].cell(row=1, column=LEADING.index(column) + 1).value,
                                     column)

    def test_a_changed_cell_is_caught(self):
        book = load_workbook(self.path)
        header, rows = self.exports["standard"]
        sheet = book["Standard"]
        sheet.cell(row=2, column=5).value = "0.5"
        self.assertEqual(len(mismatches(sheet, header, rows)), 1)


if __name__ == "__main__":
    unittest.main()
